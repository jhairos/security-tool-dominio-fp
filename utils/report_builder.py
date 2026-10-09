"""
utils/report_builder.py
=======================
Generador de los cuatro productos de evidencia:

    index.html            reporte para navegador (imagenes embebidas)
    gmail.html            tabla con una imagen por fila (data URI)
    reporte_gmail.docx    tabla con imagen por fila (copiar/pegar Gmail)
    evidencias_<X>_*.zip  todas las PNG capturadas

El titulo, la plataforma, la introduccion y el prefijo del ZIP los
define quien invoca el modulo.
"""
from __future__ import annotations

import base64
import zipfile
from datetime import datetime
from pathlib import Path
from typing import List, Optional, Tuple

from config import APP_AUTHOR
from utils.logger import get_logger

log = get_logger("utils.reports")

#: Evidencia: ``(titulo, ruta_png | None, estado)`` o, cuando se desea
#: dejar constancia del origen, ``(titulo, ruta_png, estado, referencia)``.
#: La referencia es la direccion web consultada o el comando ejecutado:
#: se imprime bajo la imagen para que el reporte sea verificable.
EvidenceItem = Tuple[str, Optional[Path], str]


def desglosar(item) -> tuple:
    """
    Normaliza una evidencia a ``(titulo, png, estado, referencia)``.

    Admite las evidencias de tres elementos de las rutinas que no
    declaran origen, de modo que ninguna rutina anterior requiere
    cambios.
    """
    if len(item) >= 4:
        titulo, png, estado, referencia = item[0], item[1], item[2], item[3]
    else:
        titulo, png, estado = item[0], item[1], item[2]
        referencia = ""
    return titulo, png, estado, (referencia or "")


def _referencia_html(referencia: str) -> str:
    """Linea de referencia bajo la imagen, en el reporte de navegador."""
    if not referencia:
        return ""
    if referencia.startswith("http"):
        destino = (f'<a href="{referencia}" style="color:#36c;'
                   f'text-decoration:none;">{referencia}</a>')
    else:
        destino = f"<code>{referencia}</code>"
    return (f'<div style="font-size:11px;color:#555;margin:6px 0 0 0;'
            f'word-break:break-all;">Referencia: {destino}</div>')

STATUS_OK = "OK"


class ReportError(RuntimeError):
    """Error controlado durante la generacion de reportes."""


# ---------------------------------------------------------------------
# Utilidades
# ---------------------------------------------------------------------
def file_to_data_uri(path: Path) -> str:
    """Convierte un PNG en data URI base64 para embeberlo en HTML."""
    raw = Path(path).read_bytes()
    enc = base64.b64encode(raw).decode("ascii")
    return f"data:image/png;base64,{enc}"


def _enlazar(texto: str) -> str:
    """Convierte las direcciones del texto en enlaces."""
    import re
    return re.sub(
        r"(https?://[^\s<>\"]+)",
        r'<a href="\1" style="color:#36c;word-break:break-all;">\1</a>',
        texto)


def intro_html(intro: str) -> str:
    """
    Compone el encabezado del reporte a partir de un texto plano.

    Los parrafos se separan con una linea en blanco y las direcciones se
    convierten en enlaces. Si no hay texto, no se emite nada: las
    rutinas que no declaran encabezado quedan igual que antes.
    """
    if not (intro or "").strip():
        return ""
    import html  # noqa: PLC0415
    parrafos = [p.strip() for p in intro.strip().split("\n\n") if p.strip()]
    # El texto incluye datos externos (respuestas DNS): se escapa antes de
    # convertir las direcciones en enlaces.
    cuerpo = "".join(
        f'<p style="margin:0 0 10px 0;">'
        f'{_enlazar(html.escape(p, quote=False)).replace(chr(10), "<br>")}</p>'
        for p in parrafos)
    return (f'<div style="margin:0 0 16px 0;font-size:13px;color:#111;'
            f'line-height:1.5;">{cuerpo}</div>')


def footer_text(author: str = APP_AUTHOR) -> str:
    return (f"Reporte generado {author} · "
            f"{datetime.now().strftime('%Y-%m-%d %H:%M')}")


# ---------------------------------------------------------------------
# ZIP de evidencias
# ---------------------------------------------------------------------
def build_evidence_zip(items: List[EvidenceItem], out_zip: Path) -> Path:
    out_zip = Path(out_zip)
    try:
        with zipfile.ZipFile(out_zip, "w",
                             compression=zipfile.ZIP_DEFLATED) as z:
            for item in items:
                _title, png_path, _status, _ref = desglosar(item)
                if png_path and Path(png_path).exists():
                    z.write(str(png_path), arcname=Path(png_path).name)
    except Exception as exc:  # noqa: BLE001
        raise ReportError(f"Error creando el ZIP de evidencias: {exc}") from exc
    log.info("ZIP de evidencias generado: %s", out_zip)
    return out_zip


# ---------------------------------------------------------------------
# HTML para navegador
# ---------------------------------------------------------------------
def build_html_navegador(report_title: str, platform: str, fecha_str: str,
                         items: List[EvidenceItem],
                         author: str = APP_AUTHOR,
                         intro: str = "") -> str:
    now = datetime.now().strftime("%Y-%m-%d %H:%M")

    parts = []
    for item in items:
        title, png_path, status, referencia = desglosar(item)
        if png_path and Path(png_path).exists():
            uri = file_to_data_uri(Path(png_path))
            parts.append(f"""
            <div style="margin:18px 0;">
              <div style="font-weight:bold;margin:0 0 6px 0;">{title}</div>
              <div style="color:#666;font-size:12px;margin:0 0 10px 0;">Estado: {status} • {Path(png_path).name}</div>
              <img src="{uri}" style="width:100%;height:auto;border:1px solid #ddd;border-radius:10px;display:block;">
              {_referencia_html(referencia)}
            </div>
            """)
        else:
            parts.append(
                f"<div style='color:#a00;margin:12px 0;'>No capturado: "
                f"<b>{title}</b> — {status}"
                f"{_referencia_html(referencia)}</div>")

    return f"""<!doctype html>
<html lang="es">
<meta charset="utf-8">
<body style="font-family:Arial;background:#fff;padding:18px;max-width:1200px;margin:auto;">
  <h2 style="margin:0 0 6px 0;">{report_title}</h2>
  <div style="color:#666;margin:0 0 14px 0;">Generado fecha• {now}</div>

  {intro_html(intro)}

  <table border="1" cellpadding="8" cellspacing="0" style="border-collapse:collapse;width:100%;">
    <tr style="background:#f2f2f2;font-weight:bold;"><th>Plataforma</th><th>Observaciones</th><th>Fecha</th><th>Escalado</th></tr>
    <tr><td>{platform}</td><td>Reporte</td><td>{fecha_str}</td><td>NA</td></tr>
  </table>

  {''.join(parts)}

  <div style="font-size:11px;color:#777;margin-top:16px;">{footer_text(author)}</div>
</body>
</html>"""


# ---------------------------------------------------------------------
# HTML para Gmail (una foto por fila)
# ---------------------------------------------------------------------
def build_html_gmail(report_title: str, platform: str, fecha_str: str,
                     items: List[EvidenceItem],
                     author: str = APP_AUTHOR,
                     intro: str = "") -> str:
    now = datetime.now().strftime("%Y-%m-%d %H:%M")

    rows = []
    for item in items:
        title, png_path, status, referencia = desglosar(item)
        if png_path and Path(png_path).exists():
            uri = file_to_data_uri(Path(png_path))
            rows.append(f"""
            <tr>
              <td style="padding:10px;border:1px solid #ddd;vertical-align:top;width:220px;">
                <div style="font-weight:bold;">{title}</div>
                <div style="color:#666;font-size:12px;">Estado: {status}</div>
              </td>
              <td style="padding:10px;border:1px solid #ddd;">
                <img src="{uri}" style="display:block;width:100%;max-width:980px;height:auto;border:1px solid #ccc;border-radius:10px;">
                {_referencia_html(referencia)}
              </td>
            </tr>
            """)
        else:
            rows.append(f"""
            <tr>
              <td style="padding:10px;border:1px solid #ddd;vertical-align:top;width:220px;">
                <div style="font-weight:bold;">{title}</div>
              </td>
              <td style="padding:10px;border:1px solid #ddd;color:#a00;">
                No capturado — {status}
                {_referencia_html(referencia)}
              </td>
            </tr>
            """)

    return f"""<!doctype html>
<html lang="es">
<meta charset="utf-8">
<body style="font-family:Arial,Helvetica,sans-serif;font-size:13px;color:#111;">
  <div style="font-size:16px;font-weight:bold;">{report_title}</div>
  <div style="color:#666;margin:6px 0 14px 0;">Generado fecha • {now}</div>

  {intro_html(intro)}

  <table cellpadding="8" cellspacing="0" border="1" style="border-collapse:collapse;border:1px solid #222;width:100%;">
    <tr style="background:#f2f2f2;font-weight:bold;">
      <td>Plataforma</td><td>Observaciones</td><td>Fecha</td><td>Escalado</td>
    </tr>
    <tr>
      <td>{platform}</td><td>Reporte</td><td>{fecha_str}</td><td>NA</td>
    </tr>
  </table>

  <div style="margin:14px 0 8px 0;font-weight:bold;">Evidencias (cada cuadro incluye la imagen)</div>

  <table cellpadding="0" cellspacing="0" border="0" style="border-collapse:collapse;width:100%;">
    {''.join(rows)}
  </table>

  <div style="font-size:11px;color:#777;margin-top:14px;">{footer_text(author)}</div>
</body>
</html>"""


# ---------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------
def build_docx(report_title: str, platform: str, fecha_str: str,
               items: List[EvidenceItem], out_docx: Path,
               image_width_inches: float = 6.6,
               author: str = APP_AUTHOR,
               intro: str = "") -> Path:
    try:
        from docx import Document  # noqa: PLC0415
        from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: PLC0415
        from docx.shared import Inches, Pt  # noqa: PLC0415
    except ImportError as exc:  # noqa: BLE001
        raise ReportError(
            "python-docx no esta instalado. Ejecute: "
            "pip install python-docx") from exc

    out_docx = Path(out_docx)
    doc = Document()

    p = doc.add_paragraph(report_title)
    p.runs[0].bold = True
    doc.add_paragraph(
        f"Generado fecha • {datetime.now().strftime('%Y-%m-%d %H:%M')}")

    # Encabezado del reporte, si la rutina lo declara.
    if (intro or "").strip():
        doc.add_paragraph("")
        for parrafo in [x.strip() for x in intro.strip().split("\n\n")
                        if x.strip()]:
            doc.add_paragraph(parrafo)
        doc.add_paragraph("")

    t = doc.add_table(rows=2, cols=4)
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    hdr[0].text = "Plataforma"
    hdr[1].text = "Observaciones"
    hdr[2].text = "Fecha"
    hdr[3].text = "Escalado"
    row = t.rows[1].cells
    row[0].text = platform
    row[1].text = "Reporte"
    row[2].text = fecha_str
    row[3].text = "NA"

    doc.add_paragraph("")
    p2 = doc.add_paragraph("Evidencias (cada cuadro incluye la imagen)")
    p2.runs[0].bold = True

    te = doc.add_table(rows=1, cols=2)
    te.style = "Table Grid"
    te.rows[0].cells[0].text = "Pantalla"
    te.rows[0].cells[1].text = "Imagen"

    img_width = Inches(image_width_inches)

    for item in items:
        title, png_path, status, referencia = desglosar(item)
        r = te.add_row().cells
        r[0].text = f"{title}\nEstado: {status}"
        if png_path and Path(png_path).exists():
            para = r[1].paragraphs[0]
            run = para.add_run()
            run.add_picture(str(png_path), width=img_width)
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        else:
            r[1].text = f"No capturado — {status}"

        # Referencia bajo la imagen: direccion consultada o comando.
        if referencia:
            pie = r[1].add_paragraph()
            corrida = pie.add_run(f"Referencia: {referencia}")
            corrida.italic = True
            corrida.font.size = Pt(8)

    doc.add_paragraph(footer_text(author))

    try:
        doc.save(str(out_docx))
    except Exception as exc:  # noqa: BLE001
        raise ReportError(f"Error generando el DOCX: {exc}") from exc

    log.info("DOCX generado: %s", out_docx)
    return out_docx


# ---------------------------------------------------------------------
# Orquestacion de los cuatro productos
# ---------------------------------------------------------------------
def build_all_reports(items: List[EvidenceItem],
                      out_dir: Path,
                      stamp: str,
                      platform: str,
                      report_title: str,
                      zip_prefix: str,
                      fecha_str: Optional[str] = None,
                      reporter=None,
                      image_width_inches: float = 6.6,
                      author: str = APP_AUTHOR,
                      intro: str = "") -> dict:
    """
    Genera los cuatro productos y devuelve
    ``{"zip":..., "html":..., "gmail_html":..., "docx":..., "errores":[...]}``.

    Un fallo en un producto no impide la generacion de los demas.
    """
    out_dir = Path(out_dir)
    fecha_str = fecha_str or datetime.now().strftime("%d-%b-%Y").upper()

    archivos: dict = {"zip": None, "html": None,
                      "gmail_html": None, "docx": None}
    errores: List[str] = []

    def _step(percent: int, message: str) -> None:
        if reporter:
            reporter.progress(percent, message)

    _step(85, "Generando ZIP de evidencias")
    try:
        archivos["zip"] = build_evidence_zip(
            items, out_dir / f"{zip_prefix}{stamp}.zip")
    except Exception as exc:  # noqa: BLE001
        log.exception("Error generando ZIP")
        errores.append(f"ZIP: {exc}")

    _step(90, "Generando index.html")
    try:
        path = out_dir / "index.html"
        path.write_text(
            build_html_navegador(report_title, platform, fecha_str, items,
                                 author, intro), encoding="utf-8")
        archivos["html"] = path
    except Exception as exc:  # noqa: BLE001
        log.exception("Error generando index.html")
        errores.append(f"index.html: {exc}")

    _step(94, "Generando gmail.html")
    try:
        path = out_dir / "gmail.html"
        path.write_text(
            build_html_gmail(report_title, platform, fecha_str, items,
                             author, intro), encoding="utf-8")
        archivos["gmail_html"] = path
    except Exception as exc:  # noqa: BLE001
        log.exception("Error generando gmail.html")
        errores.append(f"gmail.html: {exc}")

    _step(97, "Generando reporte_gmail.docx")
    try:
        archivos["docx"] = build_docx(
            report_title, platform, fecha_str, items,
            out_dir / "reporte_gmail.docx",
            image_width_inches=image_width_inches, author=author,
            intro=intro)
    except Exception as exc:  # noqa: BLE001
        log.exception("Error generando DOCX")
        errores.append(f"DOCX: {exc}")

    archivos["errores"] = errores
    return archivos
